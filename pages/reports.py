"""Octa Financial Control — Financial Reports Export."""
import streamlit as st
import io
from modules.auth import require_auth
from modules.sso import auto_login_from_url
from modules.ui_helpers import (inject_css, sidebar_nav, page_header,
                                 section_label, DARK)
from modules.database import (get_project, get_work_packages,
                               get_project_partners, get_budget_matrix,
                               get_budget_summary_by_partner, get_project_totals,
                               get_overhead_settings)
from modules.charts import (chart_by_category, chart_partner_comparison,
                             chart_spending_donut, fig_to_html)
from config import COST_CATEGORIES, CAT_LABELS, DARK as D

st.set_page_config(page_title="Reports — Octa", page_icon="📥",
                   layout="wide", initial_sidebar_state="expanded")
inject_css(); auto_login_from_url(); require_auth(); sidebar_nav()

sel_pid = st.session_state.get("selected_project_id","")
if not sel_pid: st.switch_page("app.py"); st.stop()

proj    = get_project(sel_pid)
acronym = proj.get("acronym","") or sel_pid
page_header("Financial Reports",
            f"{acronym} — Export full financial data as Excel or interactive HTML", "📥")
if st.button("← Dashboard"): st.switch_page("app.py")

muted    = D["muted"]
wps      = get_work_packages(sel_pid)
partners = get_project_partners(sel_pid)
matrix   = get_budget_matrix(sel_pid)
summaries= get_budget_summary_by_partner(sel_pid)
totals   = get_project_totals(sel_pid)
oh_pcts  = get_overhead_settings(sel_pid)
p_map    = {p["id"]: p for p in partners}
wp_map   = {w["wp_id"]: w for w in wps}


# ═══════════════════════════════════════════════════════════════════════════════
# Excel Export
# ═══════════════════════════════════════════════════════════════════════════════

def _build_excel() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    def _fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def _font(bold=False, color="1A1A2E", size=10):
        return Font(bold=bold, color=color, size=size, name="Calibri")

    def _align(h="left"):
        return Alignment(horizontal=h, vertical="center")

    thin = Side(border_style="thin", color="BBCCDD")
    def _border():
        return Border(bottom=thin)

    NAVY   = "1B2A4A"
    TEAL   = "00BCD4"
    MID    = "2E4A7A"
    LIGHT  = "EBF1F8"
    WHITE  = "FFFFFF"
    GREEN  = "6FCF97"
    AMBER  = "F6CC52"
    RED    = "FC8181"

    wb = Workbook()
    wb.remove(wb.active)

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws = wb.create_sheet("📊 Project Summary")
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 20

    row = 1
    ws.merge_cells(f"A{row}:E{row}")
    ws[f"A{row}"] = f"Financial Summary — {acronym}"
    ws[f"A{row}"].font  = _font(True, WHITE, 14)
    ws[f"A{row}"].fill  = _fill(NAVY)
    ws[f"A{row}"].alignment = _align("center")
    ws.row_dimensions[row].height = 24; row += 1

    from datetime import date
    ws.merge_cells(f"A{row}:E{row}")
    ws[f"A{row}"] = f"Generated: {date.today().isoformat()}"
    ws[f"A{row}"].font = _font(False, "888888", 9)
    row += 2

    # Headers
    for ci, h in enumerate(["Cost Category","Planned (€)","Actual Spent (€)",
                              "Remaining (€)","% Consumed"], 1):
        cell = ws.cell(row, ci, h)
        cell.font  = _font(True, WHITE)
        cell.fill  = _fill(TEAL)
        cell.alignment = _align("center")
    ws.row_dimensions[row].height = 17; row += 1

    alt = False
    for cat_key, cat_label in COST_CATEGORIES:
        cd = totals["by_category"].get(cat_key, {})
        pl = cd.get("planned",0); ac = cd.get("actual",0)
        rm = pl - ac
        pct= round(ac/pl*100,1) if pl else 0
        clr= GREEN if pct<75 else (AMBER if pct<95 else RED)

        for ci, val in enumerate([cat_label, pl, ac, rm, f"{pct}%"], 1):
            cell = ws.cell(row, ci, val)
            cell.fill = _fill("F7FAFC" if alt else "FFFFFF")
            cell.font = _font(False if ci>1 else False)
            if ci == 5:
                cell.font = _font(True, clr)
            if ci > 1 and ci < 5:
                cell.number_format = '#,##0.00 "€"'
            cell.alignment = _align("right" if ci>1 else "left")
        row += 1; alt = not alt

    # Grand total
    tp = totals["planned"]; ta = totals["actual"]
    for ci, val in enumerate(["TOTAL", tp, ta, tp-ta,
                               f"{round(ta/tp*100,1) if tp else 0}%"], 1):
        cell = ws.cell(row, ci, val)
        cell.fill  = _fill(NAVY)
        cell.font  = _font(True, WHITE, 11)
        cell.alignment = _align("right" if ci>1 else "left")
        if ci in (2,3,4): cell.number_format = '#,##0.00 "€"'
    ws.row_dimensions[row].height = 20; row += 2

    # ── One sheet per partner ─────────────────────────────────────────────────
    for ps in summaries:
        pid   = ps["partner_id"]
        pname = ps["partner_name"]
        short = (p_map.get(pid,{}).get("short_name","") or pname)[:28]
        safe  = "".join(c for c in short if c not in r'\/*?[]:')[:31]
        ws2   = wb.create_sheet(title=safe or f"P{pid}")

        ws2.column_dimensions["A"].width = 32
        ws2.column_dimensions["B"].width = 25
        ws2.column_dimensions["C"].width = 18
        ws2.column_dimensions["D"].width = 18
        ws2.column_dimensions["E"].width = 18

        r2 = 1
        ws2.merge_cells(f"A{r2}:E{r2}")
        ws2[f"A{r2}"] = f"Budget Report — {pname}"
        ws2[f"A{r2}"].font  = _font(True, WHITE, 13)
        ws2[f"A{r2}"].fill  = _fill(NAVY)
        ws2[f"A{r2}"].alignment = _align("center")
        ws2.row_dimensions[r2].height = 22; r2 += 1

        ws2.merge_cells(f"A{r2}:E{r2}")
        ws2[f"A{r2}"] = f"Project: {acronym} · Overhead: {oh_pcts.get(pid,25.0):.1f}%"
        ws2[f"A{r2}"].font = _font(False, WHITE, 10)
        ws2[f"A{r2}"].fill = _fill(MID)
        r2 += 2

        for ci, h in enumerate(["WP / Cost Category","Description",
                                  "Planned (€)","Actual (€)","Remaining (€)"], 1):
            cell2 = ws2.cell(r2, ci, h)
            cell2.font = _font(True, WHITE)
            cell2.fill = _fill(TEAL)
            cell2.alignment = _align("center")
        ws2.row_dimensions[r2].height = 17; r2 += 1

        for wp in sorted(wps, key=lambda x: x.get("wp_number","")):
            wid  = wp["wp_id"]
            wnum = wp.get("wp_number","")
            wtit = wp.get("wp_title","")

            # WP header row
            ws2.merge_cells(f"A{r2}:E{r2}")
            ws2[f"A{r2}"] = f"{wnum}: {wtit}"
            ws2[f"A{r2}"].font  = _font(True, WHITE, 10)
            ws2[f"A{r2}"].fill  = _fill(MID)
            ws2.row_dimensions[r2].height = 16; r2 += 1

            wp_data  = matrix.get(pid,{}).get(wid,{})
            alt2     = False
            wp_total_p = 0; wp_total_a = 0

            for cat_key, cat_label in COST_CATEGORIES:
                cd2 = wp_data.get(cat_key,{})
                pl2 = cd2.get("planned",0); ac2 = cd2.get("actual",0)
                rm2 = pl2 - ac2

                for ci, val in enumerate([cat_label,"",pl2,ac2,rm2], 1):
                    cell2 = ws2.cell(r2, ci, val)
                    cell2.fill = _fill("F7FAFC" if alt2 else "FFFFFF")
                    cell2.font = _font(cat_key=="overhead", "555555" if cat_key=="overhead" else "1A1A2E")
                    if ci > 2: cell2.number_format = '#,##0.00 "€"'
                    cell2.alignment = _align("right" if ci>2 else "left")
                alt2 = not alt2; r2 += 1
                wp_total_p += pl2; wp_total_a += ac2

            # WP subtotal
            for ci, val in enumerate([f"{wnum} TOTAL","",wp_total_p,wp_total_a,wp_total_p-wp_total_a], 1):
                cell2 = ws2.cell(r2, ci, val)
                cell2.fill  = _fill(TEAL)
                cell2.font  = _font(True, WHITE)
                cell2.alignment = _align("right" if ci>2 else "left")
                if ci > 2: cell2.number_format = '#,##0.00 "€"'
            ws2.row_dimensions[r2].height = 16; r2 += 2

        # Partner grand total
        tp2 = ps["planned"]; ta2 = ps["actual"]
        for ci, val in enumerate(["PARTNER TOTAL","",tp2,ta2,tp2-ta2], 1):
            cell2 = ws2.cell(r2, ci, val)
            cell2.fill  = _fill(NAVY)
            cell2.font  = _font(True, WHITE, 11)
            cell2.alignment = _align("right" if ci>2 else "left")
            if ci > 2: cell2.number_format = '#,##0.00 "€"'
        ws2.row_dimensions[r2].height = 20

        ws2.freeze_panes = "A4"
        ws2.page_setup.orientation = "landscape"
        ws2.page_setup.paperSize   = ws2.PAPERSIZE_A4
        ws2.page_setup.fitToPage   = True; ws2.page_setup.fitToWidth = 1

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return buf.read()


# ═══════════════════════════════════════════════════════════════════════════════
# Streamlit UI
# ═══════════════════════════════════════════════════════════════════════════════

section_label("📥 Excel Financial Report")
st.markdown(
    f"<p style='color:{muted};font-size:0.85rem'>"
    f"Exports a multi-sheet Excel workbook: one summary sheet + one sheet per partner "
    f"with full WP × category breakdown.</p>",
    unsafe_allow_html=True)

if st.button("🔨 Generate Excel Report", type="primary"):
    with st.spinner("Building financial report…"):
        excel = _build_excel()
    st.download_button(
        f"📥 Download {acronym}_Financial_Report.xlsx",
        data=excel,
        file_name=f"{acronym}_Financial_Report.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="dl_excel"
    )
    st.success("✅ Excel ready.")

st.markdown("---")
section_label("🌐 Interactive HTML Dashboard")
st.markdown(
    f"<p style='color:{muted};font-size:0.85rem'>"
    f"Export all financial charts as a single interactive HTML dashboard — "
    f"ready to embed in your project website.</p>",
    unsafe_allow_html=True)

if st.button("🖥️ Generate HTML Dashboard", type="primary", key="html_dash"):
    with st.spinner("Building interactive dashboard…"):
        from datetime import date as _date
        def _c(fig): return fig.to_html(full_html=False, include_plotlyjs=False)
        charts_html = "".join([
            _c(chart_by_category(totals["by_category"],"Budget by Category")),
            _c(chart_spending_donut(totals["by_category"],"Spending Mix")),
            _c(chart_partner_comparison(summaries,"Partner Budget")),
        ])
        full = f"""<!DOCTYPE html>
<html lang="en">
<head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{acronym} — Financial Dashboard</title>
<script src="https://cdn.plot.ly/plotly-latest.min.js"></script>
<style>
body{{background:#0f1421;color:#e2e8f0;font-family:Calibri,sans-serif;margin:0;padding:1rem 2rem}}
h1{{color:#00BCD4;font-size:1.8rem;margin-bottom:.3rem}}
p{{color:#8899b0;font-size:.9rem}}
.grid{{display:grid;grid-template-columns:1fr 1fr;gap:1.5rem;margin-top:1.5rem}}
.chart{{background:#1a2235;border-radius:12px;padding:1rem;border:1px solid rgba(255,255,255,.09)}}
.chart.full{{grid-column:1/-1}}
footer{{margin-top:2rem;color:#8899b0;font-size:.78rem;text-align:center}}
</style></head>
<body>
<h1>💶 {acronym} — Financial Dashboard</h1>
<p>{proj.get('proposal_title','')} · Generated {_date.today().isoformat()}</p>
<div class="grid">
  <div class="chart full">{_c(chart_by_category(totals['by_category'],'Budget by Cost Category'))}</div>
  <div class="chart">{_c(chart_spending_donut(totals['by_category'],'Spending Mix'))}</div>
  <div class="chart">{_c(chart_partner_comparison(summaries,'Partner Budget Comparison'))}</div>
</div>
<footer>Generated by Octa Financial Control · {acronym}</footer>
</body></html>"""

    st.download_button(
        f"📥 Download {acronym}_Financial_Dashboard.html",
        data=full.encode("utf-8"),
        file_name=f"{acronym}_Financial_Dashboard.html",
        mime="text/html", key="dl_html_dash"
    )
    st.success("✅ HTML dashboard ready — open in browser or embed with an <iframe>.")
